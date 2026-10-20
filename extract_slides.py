from pptx import Presentation
from pathlib import Path
import json
root=Path(r'D:\codeVisualStudioCode\dai-cuong-study-hub')
out=[]
for file in sorted((root/'slidegk').glob('*.pptx')):
    prs=Presentation(str(file))
    for num,slide in enumerate(prs.slides,1):
        texts=[]
        for shape in slide.shapes:
            if shape.has_text_frame:
                val=shape.text.strip()
                if val: texts.append(val)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        val=cell.text.strip()
                        if val: texts.append(val)
        out.append({'file':file.name,'slide':num,'text':'\n'.join(texts)})
(root/'sources').mkdir(exist_ok=True)
(root/'sources'/'slides.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
for file in sorted({x['file'] for x in out}):
    slides=[x for x in out if x['file']==file]
    print(file,len(slides),'nonempty',sum(bool(x['text']) for x in slides),'chars',sum(len(x['text']) for x in slides))
